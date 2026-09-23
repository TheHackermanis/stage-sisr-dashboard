"""Export CSV / Excel de la liste filtrée."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.tracking import STATUTS

TYPES = {"offre": "Offre publiée", "spontanee": "Candidature spontanée", "dsi_interne": "DSI interne"}

# (en-tête, fonction d'extraction)
COLONNES = [
    ("Nom", lambda d: d["nom"]),
    ("Offre", lambda d: d.get("titre") or ""),
    ("Type", lambda d: TYPES.get(d["type"], d["type"])),
    ("Contrat", lambda d: d.get("type_contrat") or ""),
    ("Statut", lambda d: STATUTS[d["statut"]]["label"]),
    ("Score", lambda d: d["score"]),
    ("Favori", lambda d: "oui" if d["favori"] else ""),
    ("Priorité", lambda d: d["priorite"] or ""),
    ("Ville", lambda d: d.get("ville") or ""),
    ("Distance (km)", lambda d: d.get("distance_km")),
    ("Secteur", lambda d: d.get("secteur") or ""),
    ("NAF", lambda d: d.get("naf") or ""),
    ("Taille", lambda d: d.get("taille_libelle") or ""),
    ("Technos", lambda d: ", ".join(d.get("technos") or [])),
    ("SIRET", lambda d: d.get("siret") or ""),
    ("Adresse", lambda d: d.get("adresse") or ""),
    ("Site web", lambda d: d.get("site_web") or ""),
    ("Lien offre", lambda d: d.get("offre_url") or ""),
    ("Page contact", lambda d: d.get("page_contact") or ""),
    ("Page recrutement", lambda d: d.get("page_recrutement") or ""),
    ("Email", lambda d: d.get("email_public") or ""),
    ("Téléphone", lambda d: d.get("tel_public") or ""),
    ("Date d'envoi", lambda d: d.get("date_envoi") or ""),
    ("Entretien", lambda d: (d.get("entretien_at") or "").replace("T", " ")),
    ("Notes", lambda d: d.get("notes") or ""),
    ("Sources", lambda d: ", ".join(d.get("sources") or [])),
]


def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    # Point-virgule + BOM UTF-8 : s'ouvre correctement dans Excel / LibreOffice en français
    w = csv.writer(buf, delimiter=";")
    w.writerow([c[0] for c in COLONNES])
    for d in rows:
        w.writerow(["" if (v := f(d)) is None else v for _, f in COLONNES])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stages SISR"
    ws.append([c[0] for c in COLONNES])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for d in rows:
        ws.append([f(d) for _, f in COLONNES])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, (titre, _) in enumerate(COLONNES, start=1):
        largeur = max([len(str(titre))] + [len(str(ws.cell(r, i).value or "")) for r in range(2, min(ws.max_row, 200) + 1)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, largeur + 2), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
